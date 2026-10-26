"""
Source workbooks are found by NAME and read by their DATA, never by Excel file
type (2026-10-26): the same workbook can be .xlsx, .xlsm or .xlsb. Triggered by
the Painting Weekly Plan arriving as .xlsb, which the Painting pipeline (it only
looked for *.xlsx) could not find.
"""

from types import SimpleNamespace

import openpyxl

from excel_compat import excel_engine_for, glob_excel, load_workbook_any
from painting.normalize import to_date as painting_to_date
from painting.reader import read_workbook as read_painting_workbook
from packing.normalize import to_date as packing_to_date


def test_glob_excel_ignores_the_extension(tmp_path):
    for name in ["DPR_Oct.xlsb", "dpr_old.XLSX", "DPR_macro.xlsm", "~$DPR_lock.xlsb",
                 "DPR_notes.txt", "Other.xlsb"]:
        (tmp_path / name).write_bytes(b"")

    expected = {"DPR_Oct.xlsb", "dpr_old.XLSX", "DPR_macro.xlsm"}
    assert {p.name for p in glob_excel(tmp_path, "*DPR*")} == expected
    # a configured extension is ignored, not required
    assert {p.name for p in glob_excel(tmp_path, "*DPR*.xlsx")} == expected
    assert {p.name for p in glob_excel(tmp_path, "*DPR*.xlsb")} == expected
    assert glob_excel(tmp_path / "missing", "*") == []


def test_engine_follows_the_file():
    assert excel_engine_for("a.xlsb") == "pyxlsb"
    assert excel_engine_for("a.XLSB") == "pyxlsb"
    assert excel_engine_for("a.xlsx") == "openpyxl"
    assert excel_engine_for("a.xlsm") == "openpyxl"


def test_painting_reads_any_sheet_name_with_the_data(tmp_path):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Weekly"          # not "Spool List"
    sheet.append(["Project Code", "Drawing No", "Spool No", "QC RFP"])
    sheet.append(["P1", "D1", "S1", None])
    path = tmp_path / "Painting Weekly Plan.xlsx"
    workbook.save(path)

    rows = read_painting_workbook(path)
    assert [(r["project_code"], r["drawing_no"], r["spool_no"]) for r in rows] == [("P1", "D1", "S1")]


def test_xlsb_is_read_into_the_same_shape(tmp_path, monkeypatch):
    class Rows:
        def __init__(self, data): self.data = data
        def __enter__(self): return self
        def __exit__(self, *a): return False
        def rows(self): return ([SimpleNamespace(v=v) for v in row] for row in self.data)

    class Book:
        sheets = ["Data"]
        def __enter__(self): return self
        def __exit__(self, *a): return False
        def get_sheet(self, name):
            return Rows([["Project Code", "Drawing No", "Spool No", "Primer Coat Date"],
                         ["P1", "D1", 12.0, 46080.0]])

    monkeypatch.setattr("pyxlsb.open_workbook", lambda path: Book())
    path = tmp_path / "Painting Weekly Plan.xlsb"
    path.write_bytes(b"")

    sheet = load_workbook_any(path).worksheets[0]
    assert [c.value for c in sheet[2]] == ["P1", "D1", 12, 46080]   # 12.0 -> 12, like openpyxl
    rows = read_painting_workbook(path)
    assert rows[0]["spool_no"] == "12" and rows[0]["primer_date"] == "2026-02-27"


def test_excel_serial_numbers_are_dates_in_date_columns():
    assert painting_to_date(46080) == "2026-02-27"
    assert packing_to_date(46080.0) == "2026-02-27"
    assert painting_to_date(12) is None
